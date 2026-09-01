from __future__ import annotations

import re
import shutil
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = Path(r"C:\Users\13423\Desktop\附件2：第八届中国研究生人工智能创新大赛项目文档模版.docx")
OUTPUT_DIR = ROOT / "submission" / "final_word"
PROJECT_NAME = "灵枢 SoulDrive"
PROJECT_FULL_NAME = "灵枢 SoulDrive：端侧多模态项目知识工作台"
TEAM_PLACEHOLDER = "待按初赛报名信息填写"
GROUP_PLACEHOLDER = "待按初赛报名信息填写"
COVER_TEAM_PLACEHOLDER = "待填团队"
COVER_GROUP_PLACEHOLDER = "待填组别"
DATE = "2026.09.01"
VERSION = "V1.0（参赛初稿）"
GIT_HISTORY_ROWS = [
    ["1", "建立端侧项目知识工作台基础", "V0.1", "项目组", "2026.06.14", "建立 SoulDrive 基础工程。"],
    ["2", "文档知识接口与本机安全控制", "V0.2", "项目组", "2026.06.22", "增加文档语义接口、知识引擎界面与本机安全控制。"],
    ["3", "思维导图生成与交互完善", "V0.3", "项目组", "2026.06.25", "完善问答结果生成导图、解析与交互。"],
    ["4", "私有检索与本机工作区加固", "V0.4", "项目组", "2026.06.29", "优化私有检索层、图谱/向量存储与安全测试。"],
    ["5", "RAG 上下文与业务产物链路优化", "V0.5", "项目组", "2026.08.15", "优化上下文、结构化解析、检索评测与 PPT 生成校验。"],
    ["6", "多模态、日程与证据查看完善", "V0.6", "项目组", "2026.09.01", "日程工作区级管理；增强多模态诊断、GraphRAG 构建与证据查看。"],
]

INTRO_TEXT = (
    "灵枢 SoulDrive 是面向 Windows 本机的端侧多模态项目知识工作台，围绕论文精读、会议纪要与项目知识管理，"
    "提供项目空间、项目对话和主对话 RAG。用户可将文档导入向量知识库，选定单篇或多篇上下文获得受限、可追溯的回答；"
    "也可授权窗口或屏幕并框选重点区域，结合语音或文字进行多模态问答。系统还支持思维导图、证据与审计链、Skills 管理，"
    "以及按技术方案、论文研究、政策项目或商务汇报生成不同风格 PPT；邮件和日程在确认后执行。默认以本地工作区和端侧模型保护数据，"
    "也可按需接入云端模型增强能力。"
)


def set_run_font(run, *, size: float = 10.5, bold: bool = False, color: str | None = None, name: str = "宋体"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), name)


def set_paragraph_border(paragraph, *, color: str = "B7B7B7", size: str = "8", space: str = "4"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for side_name in ("top", "left", "bottom", "right"):
        side = OxmlElement(f"w:{side_name}")
        side.set(qn("w:val"), "single")
        side.set(qn("w:sz"), size)
        side.set(qn("w:space"), space)
        side.set(qn("w:color"), color)
        p_bdr.append(side)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, cm: float):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "7F8C8D", size: str = "6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{edge_name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)
        borders.append(edge)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def clear_paragraph(paragraph):
    paragraph.clear()
    paragraph.paragraph_format.left_indent = None
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def add_field(paragraph, instruction: str, fallback: str, *, size: float = 10.5):
    begin = paragraph.add_run()
    set_run_font(begin, size=size)
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    fld_char.set(qn("w:dirty"), "true")
    begin._r.append(fld_char)

    instr = paragraph.add_run()
    set_run_font(instr, size=size)
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    instr._r.append(instr_text)

    separate = paragraph.add_run()
    set_run_font(separate, size=size)
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    separate._r.append(separator)

    result = paragraph.add_run(fallback)
    set_run_font(result, size=size)

    end = paragraph.add_run()
    set_run_font(end, size=size)
    end_char = OxmlElement("w:fldChar")
    end_char.set(qn("w:fldCharType"), "end")
    end._r.append(end_char)


def set_update_fields(doc):
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def ensure_heading_styles(doc):
    """The supplied template is manually styled and has no built-in Heading styles."""
    for name, level in (("Heading 1", 0), ("Heading 2", 1), ("Heading 3", 2)):
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = doc.styles["Normal"]
        p_pr = style.element.get_or_add_pPr()
        old = p_pr.find(qn("w:outlineLvl"))
        if old is not None:
            p_pr.remove(old)
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), str(level))
        p_pr.append(outline)


def replace_in_runs(paragraph, replacements: dict[str, str]):
    for run in paragraph.runs:
        for old, new in replacements.items():
            if old in run.text:
                run.text = run.text.replace(old, new)
    if not paragraph.runs and paragraph.text:
        for old, new in replacements.items():
            if old in paragraph.text:
                paragraph.text = paragraph.text.replace(old, new)


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


def set_page_number_start(section, start: int):
    sect_pr = section._sectPr
    old = sect_pr.find(qn("w:pgNumType"))
    if old is not None:
        sect_pr.remove(old)
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:start"), str(start))
    sect_pr.append(pg_num)


def make_body_paragraph(doc, text: str = "", *, indent: bool = True, before: float = 0, after: float = 6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = p.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if indent:
        fmt.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    return p


def make_list_item(doc, label: str, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = p.paragraph_format
    fmt.line_spacing = 1.45
    fmt.left_indent = Cm(0.74)
    fmt.first_line_indent = Cm(-0.74)
    fmt.space_after = Pt(4)
    lead = p.add_run(label)
    set_run_font(lead, size=10.5, bold=True)
    content = p.add_run(text)
    set_run_font(content, size=10.5)
    return p


def make_h1(doc, text: str):
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.space_before = Pt(16)
    fmt.space_after = Pt(10)
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=16, bold=True)
    return p


def make_h2(doc, text: str):
    p = doc.add_paragraph(style="Heading 2")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.space_before = Pt(12)
    fmt.space_after = Pt(8)
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=13.5, bold=True)
    return p


def make_h3(doc, text: str):
    p = doc.add_paragraph(style="Heading 3")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.space_before = Pt(8)
    fmt.space_after = Pt(5)
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=11.5, bold=True)
    return p


def add_data_table(doc, headers: list[str], rows: list[list[str]], widths: list[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, heading in enumerate(headers):
        cell = header.cells[i]
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, "E8EEF2")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        clear_paragraph(p)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(heading)
        set_run_font(run, size=9.5, bold=True)
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            set_cell_width(cell, widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            clear_paragraph(p)
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(value)
            set_run_font(run, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_architecture_block(doc):
    table = doc.add_table(rows=5, cols=1)
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    labels = [
        "用户工作入口：项目空间｜项目对话（论文精读、会议纪要）｜主对话",
        "多模态输入：文档导入｜文字｜语音｜授权窗口/屏幕｜框选重点区域",
        "知识与推理：向量知识库｜指定上下文 RAG｜本地/云端模型路由｜Skills",
        "工作产出：精读结论｜会议纪要｜思维导图｜多风格 PPT｜邮件｜日程",
        "可控与可追溯：来源证据｜操作审计｜确认门｜本机工作区",
    ]
    fills = ["E8EEF2", "F4F7F8", "E8EEF2", "F4F7F8", "E8EEF2"]
    for index, label in enumerate(labels):
        cell = table.cell(index, 0)
        set_cell_width(cell, 15.2)
        set_cell_shading(cell, fills[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        clear_paragraph(p)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(label)
        set_run_font(run, size=9.8, bold=index in (0, 2, 4))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    run = caption.add_run("图 1  SoulDrive 的端侧多模态项目知识工作闭环")
    set_run_font(run, size=9.5, color="555555")


def add_screenshot_placeholder(doc, code: str, title: str, content: str, capture: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_paragraph_border(p, color="9AA5AE", size="8", space="6")
    lead = p.add_run(f"【{code} 截图占位】{title}\n")
    set_run_font(lead, size=10.2, bold=True, color="2F4F5F")
    details = p.add_run(f"需呈现：{content}\n截取要求：{capture}")
    set_run_font(details, size=9.5, color="4E4E4E")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(6)
    run = caption.add_run(f"图 {code[1:]}  {title}（后续替换为脱敏实机截图）")
    set_run_font(run, size=9.2, color="666666")


def patch_toc_field(docx_path: Path):
    """Replace the template's static TOC SDT with a Word-updatable TOC field."""
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ns = {"w": namespace}
    with zipfile.ZipFile(docx_path, "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}
    root = etree.fromstring(entries["word/document.xml"])
    sdt = root.find(".//w:sdt", namespaces=ns)
    if sdt is None:
        raise RuntimeError("模板中未找到目录内容控件")
    content = sdt.find("w:sdtContent", namespaces=ns)
    if content is None:
        raise RuntimeError("模板目录内容控件缺少内容")
    for child in list(content):
        content.remove(child)
    p = etree.SubElement(content, f"{{{namespace}}}p")
    p_pr = etree.SubElement(p, f"{{{namespace}}}pPr")
    spacing = etree.SubElement(p_pr, f"{{{namespace}}}spacing")
    spacing.set(f"{{{namespace}}}line", "400")
    spacing.set(f"{{{namespace}}}lineRule", "auto")

    def field_run(kind: str, text: str | None = None):
        run = etree.SubElement(p, f"{{{namespace}}}r")
        r_pr = etree.SubElement(run, f"{{{namespace}}}rPr")
        fonts = etree.SubElement(r_pr, f"{{{namespace}}}rFonts")
        fonts.set(f"{{{namespace}}}ascii", "Times New Roman")
        fonts.set(f"{{{namespace}}}hAnsi", "Times New Roman")
        fonts.set(f"{{{namespace}}}eastAsia", "宋体")
        size = etree.SubElement(r_pr, f"{{{namespace}}}sz")
        size.set(f"{{{namespace}}}val", "21")
        if kind == "begin" or kind == "separate" or kind == "end":
            field = etree.SubElement(run, f"{{{namespace}}}fldChar")
            field.set(f"{{{namespace}}}fldCharType", kind)
            if kind == "begin":
                field.set(f"{{{namespace}}}dirty", "true")
        elif kind == "instr":
            instr = etree.SubElement(run, f"{{{namespace}}}instrText")
            instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            instr.text = ' TOC \\o "1-2" \\h \\z \\u '
        else:
            node = etree.SubElement(run, f"{{{namespace}}}t")
            node.text = text or "打开 Word 后自动更新目录；如未更新，请按 Ctrl+A 后按 F9。"

    field_run("begin")
    field_run("instr")
    field_run("separate")
    field_run("result")
    field_run("end")
    entries["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as temp:
        temp_path = Path(temp.name)
    try:
        with zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
            for name, data in entries.items():
                target.writestr(name, data)
        shutil.move(temp_path, docx_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def patch_cover_placeholders(docx_path: Path):
    """The cover fields live in drawing text boxes, outside python-docx paragraphs."""
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ns = {"w": namespace}
    replacements = {
        "[项目名称]": PROJECT_NAME,
        "[版本号码]": VERSION,
        "[团队名称]": COVER_TEAM_PLACEHOLDER,
        "[参赛组别]": COVER_GROUP_PLACEHOLDER,
    }
    with zipfile.ZipFile(docx_path, "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}
    root = etree.fromstring(entries["word/document.xml"])
    changed = set()
    for paragraph in root.xpath(".//w:p", namespaces=ns):
        nodes = paragraph.xpath(".//w:t", namespaces=ns)
        text = "".join(node.text or "" for node in nodes)
        if text not in replacements:
            continue
        target = replacements[text]
        target_node = next((node for node in nodes if (node.text or "") not in {"[", "]"}), None)
        if target_node is None:
            raise RuntimeError(f"封面占位符无法替换：{text}")
        for node in nodes:
            node.text = target if node is target_node else ""
        changed.add(text)
    missing = set(replacements) - changed
    if missing:
        raise RuntimeError(f"封面占位符未找到：{', '.join(sorted(missing))}")
    entries["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as temp:
        temp_path = Path(temp.name)
    try:
        with zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
            for name, data in entries.items():
                target.writestr(name, data)
        shutil.move(temp_path, docx_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def build_intro(output_path: Path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.6)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    header = section.header.paragraphs[0]
    clear_paragraph(header)
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header.add_run(PROJECT_NAME)
    set_run_font(header_run, size=9.5, color="444444")
    set_paragraph_border(header, color="000000", size="8", space="1")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("第八届中国研究生人工智能创新大赛")
    set_run_font(run, size=20, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    run = p.add_run("参赛作品简介")
    set_run_font(run, size=17, bold=True)

    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(20)
    set_paragraph_border(line, color="000000", size="8", space="1")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    run = p.add_run(PROJECT_FULL_NAME)
    set_run_font(run, size=16, bold=True)

    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.LEFT
    label.paragraph_format.space_before = Pt(4)
    label.paragraph_format.space_after = Pt(7)
    run = label.add_run("作品简介")
    set_run_font(run, size=13.5, bold=True)

    body = doc.add_paragraph()
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.first_line_indent = Cm(0.74)
    body.paragraph_format.line_spacing = 1.65
    body.paragraph_format.space_after = Pt(0)
    run = body.add_run(INTRO_TEXT)
    set_run_font(run, size=11.5)

    footer = section.footer.paragraphs[0]
    clear_paragraph(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(PROJECT_NAME)
    set_run_font(footer_run, size=9, color="666666")

    props = doc.core_properties
    props.title = "参赛作品简介"
    props.subject = PROJECT_FULL_NAME
    props.author = ""
    props.keywords = ""
    props.comments = ""
    set_update_fields(doc)
    doc.save(output_path)


def build_project_document(output_path: Path):
    shutil.copy2(TEMPLATE, output_path)
    doc = Document(output_path)

    replacements = {
        "[项目名称]": PROJECT_NAME,
        "[版本号码]": VERSION,
        "[YYYY.MM.DD]": DATE,
        "[团队名称]": TEAM_PLACEHOLDER,
        "[参赛组别]": GROUP_PLACEHOLDER,
    }
    for paragraph in doc.paragraphs:
        replace_in_runs(paragraph, replacements)
    for section in doc.sections:
        section.header.is_linked_to_previous = False
        replace_in_runs(section.header.paragraphs[0], {"项目名称": PROJECT_NAME})
    ensure_heading_styles(doc)

    history = doc.tables[0]
    for history_row, values in zip(history.rows[1:], GIT_HISTORY_ROWS):
        for column, (cell, value) in enumerate(zip(history_row.cells, values)):
            p = cell.paragraphs[0]
            clear_paragraph(p)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if column in (1, 5) else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(value)
            set_run_font(run, size=8.3)

    # Paragraph 45 is the first example-body title in the retained template.
    for paragraph in list(doc.paragraphs)[45:]:
        remove_paragraph(paragraph)

    body_section = doc.sections[3]
    body_section.footer.is_linked_to_previous = False
    footer = body_section.footer.paragraphs[0]
    clear_paragraph(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, " PAGE ", "1", size=9.5)
    set_page_number_start(body_section, 1)

    # Keep the template page/section system, while assigning actual outline levels for the updateable TOC.
    make_h1(doc, "1  项目概况")
    make_h2(doc, "1.1  背景和基础")
    make_body_paragraph(
        doc,
        "科研、项目与知识型工作常同时面对论文、会议纪要、方案文档、演示材料和沟通事项。资料分散在多个文件与会话中，"
        "使用者既需要快速检索，也需要把可用结论安全地转化为汇报、沟通和计划。传统单轮问答难以承接这种持续的项目工作流，"
        "且屏幕上的重点内容、口头补充和文档上下文往往彼此脱节。",
    )
    make_body_paragraph(
        doc,
        "灵枢 SoulDrive 的定位不是单一聊天窗口，而是运行在 Windows 本机的端侧多模态项目知识工作台。系统以项目空间为组织单元，"
        "在项目对话中服务论文精读与会议纪要等持续性业务，并提供主对话承接文档 RAG、多模态问答和跨功能调度。"
        "项目资料、选择的上下文、操作痕迹和产物均在本机工作区内组织，以便后续回看与复用。",
    )
    make_h2(doc, "1.2  场景和价值")
    make_list_item(doc, "（1）论文精读与项目沉淀。", "将论文、方案或其他资料导入知识库，在项目范围内围绕方法、结论、局限与关联问题持续追问；可指定单篇或多篇上下文，避免回答无边界扩展。")
    make_list_item(doc, "（2）会议纪要与协同准备。", "围绕会议内容形成结构化纪要、待办与后续沟通草稿，并将其与同一项目的资料、对话和产物关联，减少重复梳理。")
    make_list_item(doc, "（3）受授权的多模态问答。", "用户明确授权窗口或屏幕，并框选当前重点区域后，可结合语音或文字提问；系统只在授权范围内组织本轮视觉、文字与项目知识。")
    make_list_item(doc, "（4）知识到业务产物。", "在可选择证据范围的前提下，生成思维导图、不同汇报场景的 PPT、邮件草稿与日程草稿；对邮件发送等外部动作保留确认步骤。")
    make_body_paragraph(
        doc,
        "这一工作流的价值在于把“查资料、理解资料、组织材料、形成产物”置于同一项目空间：回答不只给出自然语言文本，"
        "还保留可定位的来源、选区或操作记录；当本地模型已满足隐私或离线要求时可完全在端侧运行，"
        "当使用者主动选择更强能力时也可按需接入兼容接口的云端模型。",
    )
    make_h2(doc, "1.3  所需支持")
    make_body_paragraph(
        doc,
        "当前原型面向 Windows 本机单用户环境。基础运行条件包括本地工作区、可用的端侧模型或用户自行配置的兼容云端模型、"
        "文档存储空间与常规网络连接（仅在云端模型或在线获取 Skills 时需要）。邮件发送需要由使用者完成 SMTP 配置并在发送前确认；"
        "日程以本机草稿、提醒和导出流程为主。",
    )
    make_body_paragraph(
        doc,
        "为保证展示公平与数据安全，演示资料应使用公开、授权或已脱敏的论文、纪要与界面；截图中不得出现个人联系方式、"
        "密钥、真实账号、绝对路径或其他敏感信息。团队名称、成员排序和参赛组别仅按报名信息补充，不在本文档中展开其他身份信息。",
    )

    doc.add_page_break()
    make_h1(doc, "2  项目规划")
    make_h2(doc, "2.1  整体目标")
    make_body_paragraph(
        doc,
        "参赛阶段的目标是完成一个可演示、可编辑、可追溯的端侧多模态项目知识工作台原型：使用者能够在项目空间内导入资料，"
        "围绕论文精读或会议纪要开展对话，在主对话中调用受范围约束的 RAG；必要时通过授权窗口/屏幕、框选区域和语音补足当下上下文；"
        "随后将已确认的内容转化为 PPT、邮件、日程或思维导图，并在审计与证据视图中回看过程。",
    )
    add_architecture_block(doc)
    make_body_paragraph(
        doc,
        "系统的总体设计遵循“本地优先、显式授权、范围受控、产物可追溯”的原则。默认链路不要求把项目资料上传到云端；"
        "云端模型是可选增强通道，而非默认依赖。对于外部动作和第三方 Skills，系统采用下载、审阅、确认启用或确认发送等步骤，"
        "将功能扩展与不可逆操作放入可见的控制边界。",
    )
    make_h2(doc, "2.2  技术创新点")
    make_list_item(doc, "（1）项目空间驱动的双对话组织。", "以项目空间承载资料和持续会话：项目对话服务论文精读、会议纪要等业务；主对话负责 RAG、多模态输入与工作台能力编排，避免把所有任务混为一个无范围聊天记录。")
    make_list_item(doc, "（2）受授权选区的多模态上下文。", "将文字、语音、授权窗口/屏幕与用户框选的重点区域组合为本轮输入，强调“用户指定什么，就理解什么”，而非后台采集全部屏幕。")
    make_list_item(doc, "（3）可选择上下文的证据式 RAG。", "文档进入向量知识库后，使用者可查看、删除、选择单篇或多篇上下文；回答围绕选定证据范围生成，并配合来源与审计信息降低不可回溯性。")
    make_list_item(doc, "（4）Skills 驱动的业务产物生成。", "将对话 Skills 与 PPT Skills 纳入统一管理；在线获取的 Skill 先下载和审核，再由使用者确认启用。PPT 根据技术方案、论文研究、政策项目或商务汇报等场景选择不同的风格能力。")
    make_list_item(doc, "（5）端侧与云端的可切换能力层。", "本地工作区与端侧模型优先满足隐私、可控与离线使用；当用户接受外部服务时，可配置兼容云端模型，以在能力与数据边界之间做自主选择。")

    doc.add_page_break()
    make_h1(doc, "3  实施方案")
    make_h2(doc, "3.1  技术可行性分析")
    make_h3(doc, "3.1.1  本机工作台与知识库基础")
    make_body_paragraph(
        doc,
        "系统面向 Windows 本机运行，项目工作区分别组织资料、索引、产物、日程、Skills 与审计记录。文档导入后可被解析并进入向量知识库，"
        "RAG 在检索到的资料与用户选择的上下文范围内组织回答。查看和删除入口使使用者能够管理已导入资料，"
        "从而避免知识库成为不可见、不可维护的黑箱。",
    )
    make_h3(doc, "3.1.2  多模态输入的权限与边界")
    make_body_paragraph(
        doc,
        "屏幕理解不采用后台全量录屏。使用者先授权当前窗口或屏幕，再以框选方式指明重点区域；文字、语音转写、OCR 或视觉理解结果"
        "只作为本轮问题的补充上下文。对于可读取的文字区域，系统可优先利用结构化文字；对于图表、图标或复杂版面，再使用视觉理解路径。"
        "这种分层方式兼顾响应速度与对授权范围的尊重。",
    )
    make_h3(doc, "3.1.3  业务产物与安全确认")
    make_body_paragraph(
        doc,
        "PPT 服务根据已选择资料生成可编辑演示文稿，并允许按汇报场景选择不同的 Skills；邮件可在完成本机 SMTP 配置、确认收件人和正文后真实发送；"
        "日程支持本机草稿、提前提醒和确认后的导出。第三方 Skills 采用隔离下载、人工查看适配信息、确认启用的流程，"
        "避免将网络来源的指令直接当作可执行操作。",
    )
    make_h3(doc, "3.1.4  模型服务的可选扩展")
    make_body_paragraph(
        doc,
        "本地模型可在不依赖云端推理的条件下完成工作台的基础服务；云端模型配置遵循用户主动填写接口、模型与密钥的方式，"
        "用于对性能、隐私和成本作选择。轻量离线模型与可移动介质的即插即用运行设想属于后续优化方向，当前不作为已实现功能或竞赛结论。",
    )

    make_h2(doc, "3.2  技术细节")
    make_h3(doc, "3.2.1  项目空间、项目对话与主对话")
    make_body_paragraph(
        doc,
        "项目空间用于承载项目名称、资料、会话、产物和本地状态。项目对话面向持续业务，例如对论文的章节、方法与结论进行精读，"
        "或围绕会议内容形成纪要；主对话则承接更通用的知识检索、能力调用与多模态问答。二者共用项目边界，但保留不同的业务入口与上下文。",
    )
    add_screenshot_placeholder(
        doc,
        "S1",
        "项目空间与项目对话",
        "项目列表、进入后的项目资料区，以及一段不含个人信息的论文精读或会议纪要对话。",
        "截取完整工作台主区域；项目名称可使用演示名称；隐藏真实文件路径、账号头像、联系人和通知内容。",
    )
    make_h3(doc, "3.2.2  文档导入、向量知识库与受范围约束的 RAG")
    make_body_paragraph(
        doc,
        "资料导入后形成可检索的向量知识库。使用者可在知识库中查看条目、管理或删除资料，并在提问前选择一篇或多篇上下文。"
        "RAG 以选定资料和检索结果为依据组织回答；当证据不足时，应提示资料范围不足，而不把模型补全内容当作事实。"
        "回答与来源定位、审计信息共同构成可回看的知识使用记录。",
    )
    add_screenshot_placeholder(
        doc,
        "S2",
        "向量知识库与指定上下文 RAG",
        "文档导入/列表、勾选单篇或多篇上下文，以及回答中的来源或证据区域。",
        "使用公开论文或自制脱敏材料；展示“范围已选择”的状态，不展示本机绝对路径、文件夹树、密钥或真实联系人。",
    )
    make_h3(doc, "3.2.3  授权屏幕、框选区域、语音与文字的多模态问答")
    make_body_paragraph(
        doc,
        "多模态功能由用户显式触发。用户可授权窗口或屏幕、截取当前画面并框选重点，再通过语音或文字提出问题。"
        "系统将选区与当轮输入结合，必要时补充 OCR 或视觉理解结果，并把视觉/语音信息作为可区分的证据来源。"
        "这一设计服务于当前工作界面的理解，不将未授权内容纳入上下文。",
    )
    add_screenshot_placeholder(
        doc,
        "S3",
        "授权屏幕、重点框选与语音/文字输入",
        "已授权的演示窗口、明显的选区边框、语音转写或文字提问，以及多模态回答的证据提示。",
        "仅使用公开网页、演示页面或自制图片；选区应覆盖论文图表、PPT 页面或文档段落等可说明能力的内容；遮蔽浏览器账号、系统托盘和个人信息。",
    )
    make_h3(doc, "3.2.4  Skills 管理与面向场景的 PPT 生成")
    make_body_paragraph(
        doc,
        "Skills 是工作台的可管理能力单元。系统支持从网络来源获取 Skill、在本机查看待审核内容、确认启用、刷新列表及删除已安装项；"
        "对话 Skills 与 PPT Skills 分别管理。PPT 生成前，使用者选择资料和汇报类型，系统可匹配技术方案、论文研究、政策项目或商务汇报等不同风格，"
        "并输出可继续编辑的演示文稿。",
    )
    add_screenshot_placeholder(
        doc,
        "S4",
        "Skills 安装、审核与管理",
        "资料处理 Skill 或 PPT Skill 的安装入口、待审核状态、启用确认与已安装列表。",
        "不要展示真实网络令牌、私有仓库地址或本机目录；优先选用公开 Skill 名称或自制演示条目。",
    )
    add_screenshot_placeholder(
        doc,
        "S5",
        "PPT 场景选择与生成结果",
        "所选资料、技术方案/论文研究/政策项目/商务汇报等场景选项、已选 PPT Skill，以及一张已生成 PPT 的预览或打开结果。",
        "PPT 内容应使用公开或脱敏材料；保留场景选择和生成成功状态，避免只放成品封面而看不到工作台能力。",
    )
    make_h3(doc, "3.2.5  邮件、日程与提醒")
    make_body_paragraph(
        doc,
        "邮件与日程模块将沟通和计划纳入工作台。邮件先形成主题、正文、收件人和附件清单草稿；冻结确认后，才允许使用已配置的本机 SMTP 发送。"
        "日程可创建标题、时间、时长、参与者与备注，支持本机提醒和确认后的日程导出。外发或导出前的确认使生成建议与实际执行保持可区分。",
    )
    add_screenshot_placeholder(
        doc,
        "S6",
        "邮件草稿、发送确认与日程提醒",
        "脱敏的邮件草稿和“确认后发送”状态；另配一张日程编辑或提醒弹窗截图。",
        "使用虚构邮箱与演示联系人；不要展示 SMTP 凭据、真实邮箱地址、实际会议链接或个人日程。",
    )
    make_h3(doc, "3.2.6  思维导图、审计与证据链")
    make_body_paragraph(
        doc,
        "系统可将选定知识组织为思维导图，帮助使用者从长文档或多段资料中把握层级关系；同时保留检索、模型路由、来源、"
        "多模态输入和产物操作等过程记录。审计视图不是用来替代人工判断，而是为“回答依据是什么、资料来自哪里、执行了何种操作”"
        "提供可回看的线索。",
    )
    add_screenshot_placeholder(
        doc,
        "S7",
        "思维导图",
        "一份论文或会议纪要在工作台中生成的节点式思维导图，并能看出主题到子节点的结构。",
        "选择公开材料；节点文字避免完整复制受版权保护的长段落，优先展示归纳后的标题和关系。",
    )
    add_screenshot_placeholder(
        doc,
        "S8",
        "证据来源与审计记录",
        "RAG 回答的来源定位、操作链或审计列表，能够显示当前问题与资料/操作之间的关联。",
        "选择演示项目；遮蔽 trace 标识、文件路径、模型密钥、个人名称和时间戳中的敏感部分。",
    )
    make_h3(doc, "3.2.7  本地优先与云端模型接入")
    make_body_paragraph(
        doc,
        "模型服务层支持本地优先和用户自选云端模型两种方式。本地链路将资料、索引、审计和常用推理尽量留在本机；"
        "云端模型通过兼容接口、默认模型和本机凭据管理进行配置，适合用户已接受外部服务且希望增强模型能力的场景。"
        "界面和文档应清楚提示当前模型模式，避免把可选云端服务误表述为默认的数据流向。",
    )
    add_screenshot_placeholder(
        doc,
        "S9",
        "本地/云端模型模式选择",
        "模型设置页或对话状态中能体现本地模式与可选云端模型配置的界面。",
        "不要截入密钥、完整接口地址、账户信息或凭据状态；可用“已配置/未配置”状态替代敏感字段。",
    )

    doc.add_page_break()
    make_h2(doc, "3.3  计划和分工")
    make_h3(doc, "3.3.1  参赛展示流程")
    make_list_item(doc, "步骤 1：进入项目空间。", "展示项目资料、项目对话以及论文精读或会议纪要中的一段典型任务。")
    make_list_item(doc, "步骤 2：限定证据范围问答。", "导入或选择文档，在向量知识库中勾选单篇/多篇上下文，通过主对话展示 RAG 回答及来源。")
    make_list_item(doc, "步骤 3：演示多模态输入。", "授权一个脱敏窗口或演示页面，框选重点区域，配合语音或文字提出问题，展示选区和回答之间的对应关系。")
    make_list_item(doc, "步骤 4：从知识到产物。", "选择 PPT 场景和相应 Skill 生成演示稿；展示邮件/日程草稿的确认流程，以及思维导图或审计视图。")
    make_list_item(doc, "步骤 5：说明模型边界。", "明确默认本地运行、云端为用户主动启用的可选增强；不把后续离线便携模型设想包装为当前已完成功能。")

    make_h3(doc, "3.3.2  交付计划")
    add_data_table(
        doc,
        ["阶段", "主要任务", "交付与验收要点"],
        [
            ["材料整理", "准备公开或脱敏的论文、会议纪要和演示文档；完善项目空间。", "资料可导入、可检索；不含账号、密钥、真实联系方式或身份信息。"],
            ["功能联调", "逐项检查项目对话、RAG、多模态、Skills、PPT、邮件/日程、思维导图和审计。", "每个功能均可在演示机复现，外部动作保留确认步骤。"],
            ["截图与视频", "按 S1--S9 补充实机截图；录制完整工作流视频。", "截图清晰、脱敏；视频突出“知识工作台”闭环，不加入未经证实的功能描述。"],
            ["提交核查", "填写报名信息字段，更新目录，核对文件名和内容边界。", "不出现报名信息之外、可能影响公平的身份信息。"],
        ],
        [2.2, 6.3, 7.0],
    )
    make_h3(doc, "3.3.3  团队分工（待按报名信息补充）")
    add_data_table(
        doc,
        ["成员", "建议负责内容", "交付物"],
        [
            ["成员 1（待填）", "项目统筹、项目空间与工作流设计", "项目文档、演示主线与材料核查"],
            ["成员 2（待填）", "RAG、向量知识库与证据/审计", "资料导入、上下文问答和来源展示"],
            ["成员 3（待填）", "多模态、Skills 与 PPT 业务模块", "授权选区/语音演示、Skill 与 PPT 场景展示"],
            ["成员 4（待填）", "客户端体验、邮件/日程及测试", "脱敏截图、日程提醒、发送确认与提交检查"],
        ],
        [3.3, 7.6, 4.6],
    )
    make_h3(doc, "3.3.4  后续优化方向")
    make_body_paragraph(
        doc,
        "后续将围绕更轻量的端侧模型、离线便携运行、更多类型资料的结构化解析、可配置的模型路由和可复现实测集继续优化。"
        "其中“小参数模型通过可移动介质即插即用”是面向便携部署的探索方向，需在安全、性能和稳定性验证后再纳入正式功能，"
        "不作为当前版本的能力承诺。",
    )

    doc.add_page_break()
    make_h1(doc, "4  参考资料")
    references = [
        "[1] SoulDrive 项目源代码、界面实现与本机测试记录，项目内部材料，2026。",
        "[2] Patrick Lewis et al. Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. NeurIPS, 2020。",
        "[3] Tauri 团队. Tauri 2 Documentation：桌面应用架构与本地运行说明。",
        "[4] 第八届中国研究生人工智能创新大赛项目文档模板与初赛作品提交规范。",
    ]
    for reference in references:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(reference)
        set_run_font(run, size=10.5)

    props = doc.core_properties
    props.title = "项目文档"
    props.subject = PROJECT_FULL_NAME
    props.author = ""
    props.keywords = ""
    props.comments = ""
    set_update_fields(doc)
    doc.save(output_path)
    patch_cover_placeholders(output_path)
    patch_toc_field(output_path)


def main():
    if not TEMPLATE.is_file():
        raise FileNotFoundError(f"未找到项目文档模板：{TEMPLATE}")
    intro_chars = len(re.sub(r"\s+", "", INTRO_TEXT))
    if intro_chars > 300:
        raise ValueError(f"作品简介正文超过 300 字：{intro_chars}")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    intro_path = OUTPUT_DIR / "待填团队_灵枢SoulDrive_参赛作品简介.docx"
    project_path = OUTPUT_DIR / "待填团队_灵枢SoulDrive_项目文档.docx"
    build_intro(intro_path)
    build_project_document(project_path)
    print(f"作品简介正文字符数（不含空白）：{intro_chars}")
    print(intro_path)
    print(project_path)


if __name__ == "__main__":
    main()
