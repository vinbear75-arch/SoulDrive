from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from build_competition_docs import GIT_HISTORY_ROWS, clear_paragraph, set_run_font


INPUT = Path(r"D:\AI_AGENTS\SoulDrive-main\submission\final_word\待填团队_灵枢SoulDrive_项目文档.docx")
OUTPUT = Path(r"D:\AI_AGENTS\SoulDrive-main\submission\word_work\项目文档_更新Git记录.docx")


def main():
    doc = Document(INPUT)
    history = doc.tables[0]
    if len(history.rows) < len(GIT_HISTORY_ROWS) + 1 or len(history.columns) != 6:
        raise RuntimeError("项目文档的记录更改历史表结构与预期不符")
    for history_row, values in zip(history.rows[1:], GIT_HISTORY_ROWS):
        for column, (cell, value) in enumerate(zip(history_row.cells, values)):
            paragraph = cell.paragraphs[0]
            clear_paragraph(paragraph)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column in (1, 5) else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.1
            run = paragraph.add_run(value)
            set_run_font(run, size=8.3)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
