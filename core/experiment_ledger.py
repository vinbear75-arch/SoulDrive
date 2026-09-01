"""Small, auditable experiment ledger built on the existing project catalog.

The ledger deliberately separates parsing, confirmation and comparison. A
parser may propose fields, but missing unit, split, dataset, time, metric
definition or evaluation script information never gets silently written.
"""

from __future__ import annotations

import csv
import io
import json
import math
import posixpath
import re
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError
from xml.etree import ElementTree

from core.project_catalog import ProjectCatalog


METRIC_ALIASES = ("metric", "metric_name", "name", "指标", "指标名")
VALUE_ALIASES = ("value", "score", "result", "数值", "值")
UNIT_ALIASES = ("unit", "metric_unit", "单位")
SPLIT_ALIASES = ("split", "data_split", "划分")
DATASET_ALIASES = ("dataset_version", "dataset", "data_version", "数据集", "数据集版本")
TIME_ALIASES = ("performed_at", "timestamp", "run_at", "date", "时间", "实验时间")
DIRECTION_ALIASES = ("direction", "higher_is_better", "方向")
DEFINITION_ALIASES = ("definition_version", "metric_definition_version", "定义版本")
SCRIPT_ALIASES = ("evaluation_script_ref", "script_version", "eval_script", "评测脚本版本")
SAMPLE_ALIASES = ("sample_count", "n", "样本数")
VARIANCE_ALIASES = ("variance", "方差")
MLFLOW_ARCHIVE_MAX_MEMBERS = 512
MLFLOW_ARCHIVE_MAX_BYTES = 16 * 1024 * 1024


def parse_metric_source(path: Path) -> dict[str, Any]:
    """Parse a local metrics source into explicit row proposals."""
    if path.is_dir():
        rows = _parse_mlflow_run_directory(path)
        return {
            "source_name": path.name,
            "source_format": "mlflow",
            "rows": rows,
            "row_count": len(rows),
        }
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        rows = _parse_delimited(path, "\t" if suffix == ".tsv" else ",")
    elif suffix == ".xlsx":
        rows = _parse_xlsx(path)
    elif suffix == ".json":
        rows = _parse_structured(path, "json")
    elif suffix in {".yaml", ".yml"}:
        rows = _parse_structured(path, "yaml")
    elif suffix in {".log", ".txt", ".md"}:
        rows = _parse_log(path)
    elif suffix == ".zip":
        rows = _parse_mlflow_archive(path)
    else:
        raise ValueError(f"unsupported metric source: {suffix or 'none'}")
    return {
        "source_name": path.name,
        "source_format": "mlflow" if suffix == ".zip" else suffix.lstrip(".") or "unknown",
        "rows": rows,
        "row_count": len(rows),
    }


def _parse_mlflow_archive(path: Path) -> list[dict[str, Any]]:
    if not zipfile.is_zipfile(path):
        raise ValueError("zip metric source is not a valid MLflow archive")
    with zipfile.ZipFile(path) as archive:
        members = _safe_mlflow_members(archive)

        def read_member(name: str) -> bytes:
            return archive.read(members[name])

        return _parse_mlflow_members(read_member, set(members))


def _parse_mlflow_run_directory(path: Path) -> list[dict[str, Any]]:
    if not path.is_dir():
        raise ValueError("MLflow run source must be a directory")
    files: dict[str, Path] = {}
    for item in path.rglob("*"):
        if not item.is_file():
            continue
        relative = item.relative_to(path).as_posix()
        files[relative] = item
    roots = _mlflow_run_roots(files)
    if len(roots) != 1:
        raise ValueError("MLflow source must contain exactly one run with meta.yaml and metrics/")
    root = roots[0]

    def read_member(name: str) -> bytes:
        item = files.get(name)
        if item is None:
            raise KeyError(name)
        return item.read_bytes()

    return _parse_mlflow_members(read_member, set(files), root=root)


def _safe_mlflow_members(archive: zipfile.ZipFile) -> dict[str, zipfile.ZipInfo]:
    infos = [info for info in archive.infolist() if not info.is_dir()]
    if len(infos) > MLFLOW_ARCHIVE_MAX_MEMBERS:
        raise ValueError("MLflow archive contains too many files")
    total_size = sum(max(int(info.file_size), 0) for info in infos)
    if total_size > MLFLOW_ARCHIVE_MAX_BYTES:
        raise ValueError("MLflow archive is too large")
    names: dict[str, zipfile.ZipInfo] = {}
    for info in infos:
        name = str(info.filename or "").replace("\\", "/").strip("/")
        parts = [part for part in name.split("/") if part]
        if not parts or any(part in {".", ".."} for part in parts):
            raise ValueError("MLflow archive contains an unsafe path")
        normalised = "/".join(parts)
        if normalised in names:
            raise ValueError("MLflow archive contains duplicate file paths")
        names[normalised] = info
    return names


def _mlflow_run_roots(files: dict[str, Any] | set[str]) -> list[str]:
    names = set(files)
    roots = []
    for name in names:
        if name == "meta.yaml":
            root = ""
        elif name.endswith("/meta.yaml"):
            root = name[: -len("/meta.yaml")]
        else:
            continue
        metrics_prefix = f"{root}/metrics/" if root else "metrics/"
        if any(item.startswith(metrics_prefix) and item != metrics_prefix for item in names):
            roots.append(root)
    return sorted(set(roots))


def _parse_mlflow_members(
    read_member: Any,
    files: set[str],
    *,
    root: str | None = None,
) -> list[dict[str, Any]]:
    roots = [root] if root is not None else _mlflow_run_roots(files)
    if len(roots) != 1:
        raise ValueError("MLflow source must contain exactly one run with meta.yaml and metrics/")
    run_root = roots[0]
    prefix = f"{run_root}/" if run_root else ""
    meta_name = f"{prefix}meta.yaml"
    meta = _load_yaml_mapping(read_member(meta_name).decode("utf-8-sig"))
    params = _read_mlflow_kv_files(read_member, files, f"{prefix}params/")
    tags = _read_mlflow_kv_files(read_member, files, f"{prefix}tags/")
    fields = _merge_mlflow_fields(params, tags)
    performed_at = _mlflow_timestamp(
        _mlflow_field(fields, "performed_at", "") or meta.get("start_time")
    )
    metric_prefix = f"{prefix}metrics/"
    metric_names = sorted(
        name for name in files
        if name.startswith(metric_prefix) and name != metric_prefix and not name.endswith("/")
    )
    if not metric_names:
        raise ValueError("MLflow run does not contain metrics")
    rows: list[dict[str, Any]] = []
    for metric_path in metric_names:
        metric_name = metric_path[len(metric_prefix):]
        selected: tuple[int, int, float, int] | None = None
        for line_number, line in enumerate(
            read_member(metric_path).decode("utf-8-sig").splitlines(), 1
        ):
            if not line.strip():
                continue
            parts = line.split()
            if len(parts) not in {2, 3, 5}:
                raise ValueError(f"malformed MLflow metric at {metric_path}:line:{line_number}")
            try:
                timestamp = int(parts[0])
                value = float(parts[1])
                step = int(parts[2]) if len(parts) >= 3 else 0
            except (TypeError, ValueError) as exc:
                raise ValueError(f"malformed MLflow metric at {metric_path}:line:{line_number}") from exc
            if not math.isfinite(value):
                raise ValueError(f"non-finite MLflow metric at {metric_path}:line:{line_number}")
            candidate = (step, timestamp, value, line_number)
            if selected is None or candidate[:3] > selected[:3]:
                selected = candidate
        if selected is None:
            raise ValueError(f"MLflow metric is empty: {metric_path}")
        _, _, value, line_number = selected
        row = {
            "metric": metric_name,
            "value": value,
            "performed_at": performed_at,
            "dataset_version": _mlflow_field(fields, "dataset_version", metric_name),
        "config_ref": _mlflow_field(fields, "config_ref", metric_name),
            "evaluation_script_ref": _mlflow_field(fields, "evaluation_script_ref", metric_name),
            "environment_ref": _mlflow_field(fields, "environment_ref", metric_name),
            "unit": _mlflow_field(fields, "unit", metric_name),
            "direction": _mlflow_field(fields, "direction", metric_name),
            "definition_version": _mlflow_field(fields, "definition_version", metric_name),
            "split": _mlflow_field(fields, "split", metric_name),
            "__source_locator": f"{metric_path}:line:{line_number}",
        }
        rows.append(row)
    return rows


def _read_mlflow_kv_files(read_member: Any, files: set[str], prefix: str) -> dict[str, str]:
    values: dict[str, str] = {}
    for name in sorted(files):
        if not name.startswith(prefix) or name.endswith("/"):
            continue
        key = name[len(prefix):]
        if not key or "/" in key:
            continue
        values[key] = read_member(name).decode("utf-8-sig").strip()
    return values


def _merge_mlflow_fields(params: dict[str, str], tags: dict[str, str]) -> dict[str, str]:
    fields = dict(params)
    for key, value in tags.items():
        if key in fields and fields[key] != value and key.startswith("souldrive."):
            raise ValueError(f"conflicting MLflow provenance field: {key}")
        fields[key] = value
    return fields


def _load_yaml_mapping(text: str) -> dict[str, Any]:
    try:
        import yaml

        value = yaml.safe_load(text)
    except ImportError:
        value = _simple_yaml(text)
    if not isinstance(value, dict):
        raise ValueError("MLflow meta.yaml must contain an object")
    return value


def _mlflow_field(fields: dict[str, str], field: str, metric_name: str) -> str:
    aliases = {
        "unit": ("unit", "metric_unit"),
        "direction": ("direction",),
        "definition_version": ("definition_version", "metric_definition_version"),
        "dataset_version": ("dataset_version",),
        "split": ("split",),
        "evaluation_script_ref": ("evaluation_script_ref", "eval_script"),
        "config_ref": ("config_ref",),
        "environment_ref": ("environment_ref",),
        "performed_at": ("performed_at",),
    }
    for alias in aliases.get(field, (field,)):
        candidates = []
        if metric_name:
            candidates.append(f"souldrive.metric.{metric_name}.{alias}")
        candidates.append(f"souldrive.{alias}")
        for key in candidates:
            if fields.get(key) not in (None, ""):
                return str(fields[key]).strip()
    return ""


def _mlflow_timestamp(value: Any) -> str:
    if value in (None, ""):
        return ""
    text = str(value).strip()
    try:
        milliseconds = int(float(text))
    except (TypeError, ValueError):
        return text
    return datetime.fromtimestamp(milliseconds / 1000.0, tz=timezone.utc).isoformat()


def plan_metric_import(
    path: Path,
    *,
    experiment_name: str,
    source_name: str | None = None,
    performed_at: str | None = None,
    dataset_version: str = "",
    config_ref: str = "",
    evaluation_script_ref: str = "",
    environment_ref: str = "",
    unit: str = "",
    direction: str = "",
    definition_version: str = "",
    split: str = "",
    confirmed: bool = False,
) -> dict[str, Any]:
    """Return a confirmation card or an explicit write-ready proposal."""
    parsed = parse_metric_source(path)
    proposals: list[dict[str, Any]] = []
    missing: list[dict[str, Any]] = []
    for index, raw in enumerate(parsed["rows"], 1):
        row = _normalise_metric_row(raw, index)
        row["unit"] = row.get("unit") or str(unit or "").strip()
        row["direction"] = _normalise_direction(row.get("direction") or direction)
        row["definition_version"] = row.get("definition_version") or str(definition_version or "").strip()
        row["split"] = row.get("split") or str(split or "").strip()
        row["dataset_version"] = row.get("dataset_version") or str(dataset_version or "").strip()
        row["performed_at"] = row.get("performed_at") or str(performed_at or "").strip()
        row["evaluation_script_ref"] = row.get("evaluation_script_ref") or str(evaluation_script_ref or "").strip()
        row["config_ref"] = row.get("config_ref") or str(config_ref or "").strip()
        row["environment_ref"] = row.get("environment_ref") or str(environment_ref or "").strip()
        row_missing = [
            field for field in (
                "metric_name", "value", "unit", "direction", "definition_version",
                "split", "dataset_version", "performed_at", "evaluation_script_ref",
            ) if row.get(field) in (None, "")
        ]
        if row.get("value") is not None and not _finite_number(row["value"]):
            row_missing.append("value")
        try:
            if row.get("performed_at"):
                row["performed_at"] = _canonical_timestamp(row["performed_at"])
        except ValueError:
            row_missing.append("performed_at")
        if row_missing:
            missing.append({"source_locator": row["source_locator"], "fields": sorted(set(row_missing))})
        else:
            row["value"] = float(row["value"])
            if row.get("sample_count") is not None:
                row["sample_count"] = _non_negative_int(row["sample_count"], "sample_count")
            if row.get("variance") is not None:
                row["variance"] = _non_negative_float(row["variance"], "variance")
            proposals.append(row)

    missing_fields = sorted({field for item in missing for field in item["fields"]})
    ready = bool(proposals) and not missing
    return {
        "schema_version": "souldrive_experiment_import_plan_v0.1",
        "source_name": str(source_name or parsed["source_name"]),
        "source_format": parsed["source_format"],
        "experiment_name": str(experiment_name or "").strip(),
        "row_count": parsed["row_count"],
        "valid_row_count": len(proposals),
        "missing": missing,
        "missing_fields": missing_fields,
        "requires_confirmation": ready and not confirmed,
        "ready": ready,
        "confirmed": bool(confirmed),
        "rows": proposals,
        "writes": bool(ready and confirmed),
    }


def import_metric_source(
    catalog: ProjectCatalog,
    *,
    project_id: str,
    evidence_version_id: str,
    path: Path,
    experiment_name: str,
    confirmed: bool = False,
    source_name: str | None = None,
    **defaults: Any,
) -> dict[str, Any]:
    """Plan first; write only after all required fields are confirmed."""
    plan = plan_metric_import(
        path,
        experiment_name=experiment_name,
        source_name=source_name,
        confirmed=confirmed,
        **defaults,
    )
    if not plan["writes"]:
        return {"status": "needs_confirmation" if plan["missing_fields"] else "confirmation_required", "plan": plan}
    project_id = str(project_id or "").strip()
    evidence_version_id = str(evidence_version_id or "").strip()
    if not project_id or not evidence_version_id:
        raise ValueError("project_id and evidence_version_id are required")
    evidence = catalog.conn.execute(
        """
        SELECT a.project_id, v.file_ref, v.content_hash
        FROM asset_versions v JOIN assets a ON a.asset_id = v.asset_id
        WHERE v.version_id = ?
        """,
        (evidence_version_id,),
    ).fetchone()
    if evidence is None:
        raise ValueError("evidence version does not exist")
    if str(evidence["project_id"]) != project_id:
        raise ValueError("evidence version belongs to another project")
    experiment_row = catalog.conn.execute(
        "SELECT experiment_id FROM experiments WHERE project_id = ? AND name = ? LIMIT 1",
        (project_id, plan["experiment_name"]),
    ).fetchone()
    experiment_id = str(experiment_row[0]) if experiment_row else catalog.create_experiment(
        project_id, plan["experiment_name"]
    )
    first = plan["rows"][0]
    existing = catalog.conn.execute(
        """
        SELECT run_id FROM experiment_runs
        WHERE experiment_id = ? AND evidence_version_id = ? AND performed_at = ?
          AND dataset_version = ? AND evaluation_script_ref = ?
          AND config_ref = ? AND environment_ref = ?
        LIMIT 1
        """,
        (
            experiment_id, evidence_version_id, first["performed_at"],
            first["dataset_version"], first["evaluation_script_ref"],
            first["config_ref"], first["environment_ref"],
        ),
    ).fetchone()
    run_id = str(existing[0]) if existing else catalog.create_experiment_run(
        experiment_id,
        first["performed_at"],
        evidence_version_id=evidence_version_id,
        dataset_version=first["dataset_version"],
        config_ref=first["config_ref"],
        evaluation_script_ref=first["evaluation_script_ref"],
        environment_ref=first["environment_ref"],
    )
    for row in plan["rows"]:
        metric_id = catalog.create_metric_definition(
            row["metric_name"], row["unit"], row["direction"], row["definition_version"]
        )
        definition = catalog.conn.execute(
            "SELECT direction FROM metric_definitions WHERE metric_id = ?", (metric_id,)
        ).fetchone()
        if definition is None or str(definition[0]) != row["direction"]:
            raise ValueError(f"metric definition conflict: {row['metric_name']}")
        catalog.record_measurement(
            run_id,
            metric_id,
            row["value"],
            row["split"],
            sample_count=row.get("sample_count"),
            variance=row.get("variance"),
            source_locator=row["source_locator"],
        )
    return {
        "status": "imported",
        "run_id": run_id,
        "project_id": project_id,
        "evidence_version_id": evidence_version_id,
        "source_file_ref": str(evidence["file_ref"]),
        "plan": plan,
    }


def previous_month_bounds(reference_at: str | None = None, timezone_name: str = "Asia/Shanghai") -> dict[str, str]:
    try:
        zone = ZoneInfo(timezone_name)
    except ZoneInfoNotFoundError as exc:
        raise ValueError(f"unknown timezone: {timezone_name}") from exc
    value = datetime.now(timezone.utc) if reference_at is None else _parse_timestamp(reference_at)
    local = value.astimezone(zone)
    current_month = local.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    previous_month = (current_month - timedelta(days=1)).replace(day=1)
    return {
        "timezone": timezone_name,
        "label": previous_month.strftime("%Y-%m"),
        "performed_from": previous_month.astimezone(timezone.utc).isoformat(),
        "performed_to": current_month.astimezone(timezone.utc).isoformat(),
    }


def find_previous_month_runs(
    catalog: ProjectCatalog,
    project_id: str,
    *,
    reference_at: str | None = None,
    timezone_name: str = "Asia/Shanghai",
    experiment_name: str | None = None,
) -> dict[str, Any]:
    bounds = previous_month_bounds(reference_at, timezone_name)
    runs = catalog.list_experiment_runs(
        project_id,
        performed_from=bounds["performed_from"],
        performed_to=bounds["performed_to"],
        experiment_name=experiment_name,
    )
    return {
        "status": "missing" if not runs else ("ready" if len(runs) == 1 else "needs_confirmation"),
        "bounds": bounds,
        "candidate_count": len(runs),
        "candidates": runs,
    }


def summarize_experiment_repeats(
    catalog: ProjectCatalog,
    run_ids: list[str],
    *,
    project_id: str | None = None,
) -> dict[str, Any]:
    """Summarize explicitly selected repeats without mixing experiments."""
    normalized_ids = list(dict.fromkeys(str(run_id or "").strip() for run_id in run_ids if str(run_id or "").strip()))[:32]
    runs = [catalog.get_experiment_run(run_id) for run_id in normalized_ids]
    missing = [run_id for run_id, run in zip(normalized_ids, runs) if run is None]
    if missing:
        return {"status": "missing", "missing_run_ids": missing, "run_ids": normalized_ids}
    loaded = [run for run in runs if run is not None]
    projects = {str(run["project_id"]) for run in loaded}
    experiments = {str(run["experiment_id"]) for run in loaded}
    if project_id and projects != {str(project_id)}:
        return {"status": "rejected", "reason": "cross_project", "run_ids": normalized_ids}
    if len(projects) != 1:
        return {"status": "rejected", "reason": "cross_project", "run_ids": normalized_ids}
    if len(experiments) != 1:
        return {"status": "rejected", "reason": "mixed_experiment", "run_ids": normalized_ids}

    grouped: dict[tuple[str, str, str, str], list[tuple[dict[str, Any], dict[str, Any]]]] = {}
    for run in loaded:
        for measurement in run["measurements"]:
            grouped.setdefault(_metric_key(measurement), []).append((run, measurement))
    metrics: list[dict[str, Any]] = []
    for key in sorted(grouped):
        items = grouped[key]
        directions = {str(measurement.get("direction") or "") for _, measurement in items}
        if len(directions) != 1:
            continue
        values = [float(measurement["value"]) for _, measurement in items]
        mean = sum(values) / len(values)
        variance = None
        confidence_interval_95 = None
        if len(values) >= 2:
            variance = sum((value - mean) ** 2 for value in values) / (len(values) - 1)
            standard_error = math.sqrt(variance / len(values))
            margin = 1.96 * standard_error
            confidence_interval_95 = [round(mean - margin, 12), round(mean + margin, 12)]
        metrics.append({
            "metric": key[0],
            "unit": key[1],
            "definition_version": key[2],
            "split": key[3],
            "direction": next(iter(directions)),
            "run_count": len(values),
            "mean": round(mean, 12),
            "variance": round(variance, 12) if variance is not None else None,
            "confidence_interval_95": confidence_interval_95,
            "confidence_method": "normal_approximation" if confidence_interval_95 else None,
            "sources": [_measurement_ref(run, measurement) for run, measurement in items],
        })
    return {
        "schema_version": "souldrive_experiment_repeat_summary_v0.1",
        "status": "ready" if metrics else "not_comparable",
        "project_id": next(iter(projects)),
        "experiment_id": next(iter(experiments)),
        "run_ids": normalized_ids,
        "run_count": len(loaded),
        "metrics": metrics,
    }


def compare_experiment_runs(
    catalog: ProjectCatalog,
    baseline_run_id: str,
    current_run_id: str,
    *,
    project_id: str | None = None,
) -> dict[str, Any]:
    baseline = catalog.get_experiment_run(baseline_run_id)
    current = catalog.get_experiment_run(current_run_id)
    if baseline is None or current is None:
        return {"status": "missing", "baseline_run_id": baseline_run_id, "current_run_id": current_run_id}
    if project_id and {str(baseline["project_id"]), str(current["project_id"])} != {str(project_id)}:
        return {"status": "rejected", "reason": "cross_project", "baseline": _run_ref(baseline), "current": _run_ref(current)}
    checks = {
        "same_project": baseline["project_id"] == current["project_id"],
        "dataset_version": _same_non_empty(baseline.get("dataset_version"), current.get("dataset_version")),
        "evaluation_script_ref": _same_non_empty(baseline.get("evaluation_script_ref"), current.get("evaluation_script_ref")),
    }
    baseline_metrics = {_metric_key(item): item for item in baseline["measurements"]}
    current_metrics = {_metric_key(item): item for item in current["measurements"]}
    candidate_changes: list[dict[str, Any]] = []
    missing_metrics: list[dict[str, Any]] = []
    for key in sorted(set(baseline_metrics) | set(current_metrics)):
        left = baseline_metrics.get(key)
        right = current_metrics.get(key)
        if left is None or right is None:
            missing_metrics.append({"metric": key[0], "unit": key[1], "definition_version": key[2], "split": key[3], "missing_in": "baseline" if left is None else "current"})
            continue
        if left["direction"] != right["direction"]:
            missing_metrics.append({"metric": key[0], "unit": key[1], "definition_version": key[2], "split": key[3], "missing_in": "incompatible_direction"})
            continue
        baseline_value = float(left["value"])
        current_value = float(right["value"])
        delta = round(current_value - baseline_value, 12)
        relative = None if baseline_value == 0 else round(delta / abs(baseline_value), 12)
        if left["direction"] == "higher":
            direction = "improved" if delta > 0 else "regressed" if delta < 0 else "unchanged"
        else:
            direction = "improved" if delta < 0 else "regressed" if delta > 0 else "unchanged"
        candidate_changes.append({
            "evidence_id": f"M{len(candidate_changes) + 1}",
            "metric": key[0],
            "unit": key[1],
            "definition_version": key[2],
            "split": key[3],
            "direction": left["direction"],
            "baseline_value": baseline_value,
            "current_value": current_value,
            "absolute_delta": delta,
            "relative_delta": relative,
            "assessment": direction,
            "baseline_source": _measurement_ref(baseline, left),
            "current_source": _measurement_ref(current, right),
        })
    comparable = bool(
        checks["same_project"]
        and all(checks[key] for key in ("dataset_version", "evaluation_script_ref"))
        and candidate_changes
        and not missing_metrics
    )
    # A partial match is not a safe basis for a percentage or absolute delta.
    # Keep the matching count for diagnostics, but expose no changes unless the
    # complete provenance and metric set passes the comparability gate.
    changes = candidate_changes if comparable else []
    factors = _candidate_factors(baseline, current)
    return {
        "schema_version": "souldrive_experiment_comparison_v0.1",
        "status": "comparable" if comparable else "not_comparable",
        "project_id": baseline["project_id"],
        "baseline": _run_ref(baseline),
        "current": _run_ref(current),
        "compatibility": {
            **checks,
            "matching_metric_count": len(candidate_changes),
            "missing_or_incompatible_metric_count": len(missing_metrics),
            "missing_or_incompatible_metrics": missing_metrics,
            "change_computation": "enabled" if comparable else "blocked",
        },
        "changes": changes,
        "factors": factors,
        "cause_policy": "差异与配置变化只能作为相关证据；没有区分实验时不得写成已证实原因。",
        "next_validation": _next_validation(factors) if comparable else "先补齐数据集、评测脚本版本和同定义指标后再比较。",
    }


def build_experiment_evidence(comparison: dict[str, Any] | None) -> list[dict[str, Any]]:
    """Render a comparable ledger result as bounded, citable ``M#`` evidence."""
    if not isinstance(comparison, dict) or comparison.get("status") != "comparable":
        return []
    compatibility = comparison.get("compatibility")
    if not isinstance(compatibility, dict) or compatibility.get("change_computation") != "enabled":
        return []
    baseline = comparison.get("baseline")
    current = comparison.get("current")
    if not isinstance(baseline, dict) or not isinstance(current, dict):
        return []

    factors = comparison.get("factors") if isinstance(comparison.get("factors"), list) else []
    factor_text = "; ".join(
        f"{_bounded_text(item.get('label') or item.get('factor'), 80)}："
        f"基线 {_bounded_text(item.get('baseline'), 160)}；当前 {_bounded_text(item.get('current'), 160)}"
        for item in factors[:4]
        if isinstance(item, dict) and _bounded_text(item.get("label") or item.get("factor"), 80)
    )
    evidence: list[dict[str, Any]] = []
    raw_changes = comparison.get("changes")
    if not isinstance(raw_changes, list):
        return []
    for raw_change in raw_changes[:8]:
        if not isinstance(raw_change, dict):
            continue
        metric = _bounded_text(raw_change.get("metric"), 120)
        values = {
            key: _finite_float(raw_change.get(key))
            for key in ("baseline_value", "current_value", "absolute_delta")
        }
        if not metric or any(value is None for value in values.values()):
            continue
        relative_delta = _finite_float(raw_change.get("relative_delta"))
        unit = _bounded_text(raw_change.get("unit"), 64)
        split = _bounded_text(raw_change.get("split"), 64)
        definition_version = _bounded_text(raw_change.get("definition_version"), 128)
        assessment = _bounded_text(raw_change.get("assessment"), 32)
        baseline_source = _normalise_public_source(raw_change.get("baseline_source"))
        current_source = _normalise_public_source(raw_change.get("current_source"))
        context_parts = [
            f"指标={metric}",
            f"基线值={_format_number(values['baseline_value'])}{unit}",
            f"当前值={_format_number(values['current_value'])}{unit}",
            f"绝对变化={_format_signed_number(values['absolute_delta'])}{unit}",
        ]
        if relative_delta is not None:
            context_parts.append(f"相对变化={_format_signed_number(relative_delta * 100)}%")
        if assessment:
            context_parts.append(f"判断={_assessment_label(assessment)}")
        if split:
            context_parts.append(f"split={split}")
        if definition_version:
            context_parts.append(f"指标定义={definition_version}")
        dataset_version = _bounded_text(baseline.get("dataset_version"), 160)
        if dataset_version:
            context_parts.append(f"数据集版本={dataset_version}")
        evaluation_script = _bounded_text(baseline.get("evaluation_script_ref"), 260)
        if evaluation_script:
            context_parts.append(f"评测脚本={evaluation_script}")
        baseline_locator = _measurement_locator_text(baseline_source, baseline)
        current_locator = _measurement_locator_text(current_source, current)
        if baseline_locator:
            context_parts.append(f"基线来源={baseline_locator}")
        if current_locator:
            context_parts.append(f"当前来源={current_locator}")
        if factor_text:
            context_parts.append(f"账本记录的候选变化因素={factor_text}")
        context_parts.append("这些候选因素只能作为相关性线索，不能单独证明因果")
        context = "；".join(context_parts) + "。"
        evidence_id = f"M{len(evidence) + 1}"
        evidence.append({
            "id": evidence_id,
            "evidence_id": evidence_id,
            "evidence_type": "experiment_measurement",
            "modality": "experiment",
            "score": 1.0,
            "source_filename": _measurement_source_filename(current_source, current) or "实验账本",
            "source_filenames": list(dict.fromkeys(
                item
                for item in (
                    _measurement_source_filename(baseline_source, baseline),
                    _measurement_source_filename(current_source, current),
                )
                if item
            )),
            "document_title": f"{_bounded_text(current.get('experiment_name'), 160) or '实验'}结构化对比",
            "section": f"实验测量：{metric}",
            "snippet": context,
            "context": context,
            "metric": metric,
            "unit": unit,
            "split": split,
            "definition_version": definition_version,
            "direction": _bounded_text(raw_change.get("direction"), 32),
            "baseline_value": values["baseline_value"],
            "current_value": values["current_value"],
            "absolute_delta": values["absolute_delta"],
            "relative_delta": relative_delta,
            "assessment": assessment,
            "baseline_source": baseline_source,
            "current_source": current_source,
            "support_method": "experiment_ledger_comparability_gate_v0.1",
        })
    return evidence


def build_experiment_evidence_text(evidence: list[dict[str, Any]] | None) -> str:
    """Render structured experiment measurements without mixing them with documents."""
    if not evidence:
        return ""
    lines = [
        "【结构化实验测量证据（M#）】",
        "以下数值来自当前项目中显式选择的两次实验运行，并已通过项目、数据集版本、评测脚本和完整指标集合可比性门；M# 只能支撑其中显示的测量变化，不能单独证明原因。",
    ]
    for item in evidence[:8]:
        evidence_id = _bounded_text(item.get("id"), 16) or "M?"
        content = _bounded_text(item.get("context") or item.get("snippet"), 1800)
        if content:
            lines.append(f"[{evidence_id} | 实验测量来源={_bounded_text(item.get('source_filename'), 260)}]\n{content}")
    return "\n".join(lines) if len(lines) > 2 else ""


def _normalise_public_source(value: Any) -> dict[str, str]:
    if not isinstance(value, dict):
        return {}
    return {
        key: _bounded_text(value.get(key), limit)
        for key, limit in (("run_id", 128), ("evidence_version_id", 128), ("file_ref", 260), ("source_locator", 260))
        if _bounded_text(value.get(key), limit)
    }


def _measurement_source_filename(source: dict[str, Any], run: dict[str, Any]) -> str:
    return (
        _bounded_text(source.get("file_ref"), 260)
        or _bounded_text(run.get("file_ref"), 260)
        or _bounded_text(run.get("evidence_title"), 260)
        or _bounded_text(run.get("run_id"), 128)
    )


def _measurement_locator_text(source: dict[str, Any], run: dict[str, Any]) -> str:
    filename = _measurement_source_filename(source, run)
    locator = _bounded_text(source.get("source_locator"), 260)
    if filename and locator:
        return f"{filename} / {locator}"
    return filename or locator


def _bounded_text(value: Any, limit: int) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\x00", " ")).strip()[:limit]


def _finite_float(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return round(number, 12) if math.isfinite(number) else None


def _format_number(value: float | None) -> str:
    return "" if value is None else f"{value:.12g}"


def _format_signed_number(value: float | None) -> str:
    if value is None:
        return ""
    return f"{value:+.12g}"


def _assessment_label(value: str) -> str:
    return {
        "improved": "改善",
        "regressed": "退化",
        "unchanged": "无变化",
    }.get(value, value)


def _parse_delimited(path: Path, delimiter: str) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle, delimiter=delimiter)]


def _parse_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return _parse_xlsx_xml_fallback(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        try:
            return _parse_xlsx_xml_fallback(path)
        except Exception as fallback_exc:
            raise ValueError(f"xlsx source could not be parsed: {exc}") from fallback_exc
    try:
        rows: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                continue
            headers = [str(value or "").strip() for value in values[0]]
            for row_index, values_row in enumerate(values[1:], 2):
                rows.append({"__source_locator": f"sheet:{sheet.title}/row:{row_index}", **{header: value for header, value in zip(headers, values_row) if header}})
        return rows
    finally:
        workbook.close()


def _parse_xlsx_xml_fallback(path: Path) -> list[dict[str, Any]]:
    """Read bounded worksheet XML when a lightweight export lacks full XLSX parts."""
    if not zipfile.is_zipfile(path):
        raise ValueError("xlsx source is not a valid zip archive")
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        worksheets = sorted(
            name for name in names
            if re.match(r"xl/worksheets/sheet\d+\.xml$", name)
        )
        if not worksheets:
            raise ValueError("xlsx source contains no worksheet XML")
        shared_strings = _xlsx_shared_strings(archive, names)
        labels = _xlsx_sheet_labels(archive, names, worksheets)
        parsed_rows: list[dict[str, Any]] = []
        for worksheet in worksheets:
            root = ElementTree.fromstring(archive.read(worksheet))
            physical_rows: list[tuple[int, list[str]]] = []
            for position, row in enumerate(
                (item for item in root.iter() if _xml_local_name(item.tag) == "row"),
                start=1,
            ):
                row_number = _xlsx_row_number(row, position)
                cells = [
                    _xlsx_cell_value(cell, shared_strings)
                    for cell in row
                    if _xml_local_name(cell.tag) == "c"
                ]
                if any(value.strip() for value in cells):
                    physical_rows.append((row_number, cells))
            if not physical_rows:
                continue

            sheet_label = labels.get(worksheet) or Path(worksheet).stem
            headers = [value.strip() for value in physical_rows[0][1]]
            if len(headers) == 1 and ("\n" in headers[0] or "\r" in headers[0]):
                embedded = list(csv.reader(io.StringIO(headers[0])))
                if len(embedded) > 1:
                    embedded_headers = [str(value or "").strip() for value in embedded[0]]
                    for row_number, values in enumerate(embedded[1:], 2):
                        parsed_rows.append({
                            "__source_locator": f"sheet:{sheet_label}/embedded_row:{row_number}",
                            **{
                                header: value
                                for header, value in zip(embedded_headers, values)
                                if header
                            },
                        })
                    continue

            for row_number, values in physical_rows[1:]:
                parsed_rows.append({
                    "__source_locator": f"sheet:{sheet_label}/row:{row_number}",
                    **{
                        header: value
                        for header, value in zip(headers, values)
                        if header
                    },
                })
        return [row for row in parsed_rows if any(key != "__source_locator" for key in row)]


def _xlsx_shared_strings(archive: zipfile.ZipFile, names: set[str]) -> list[str]:
    if "xl/sharedStrings.xml" not in names:
        return []
    root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
    return [
        "".join(
            str(item.text or "")
            for item in value.iter()
            if _xml_local_name(item.tag) == "t"
        )
        for value in root.iter()
        if _xml_local_name(value.tag) == "si"
    ]


def _xlsx_sheet_labels(
    archive: zipfile.ZipFile,
    names: set[str],
    worksheets: list[str],
) -> dict[str, str]:
    if "xl/workbook.xml" not in names or "xl/_rels/workbook.xml.rels" not in names:
        return {}
    workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        item.attrib.get("Id"): item.attrib.get("Target")
        for item in relationships
        if item.attrib.get("Id") and item.attrib.get("Target")
    }
    labels: dict[str, str] = {}
    for sheet in workbook.iter():
        if _xml_local_name(sheet.tag) != "sheet":
            continue
        relationship_id = next(
            (value for key, value in sheet.attrib.items() if key == "id" or key.endswith("}id")),
            None,
        )
        target = targets.get(relationship_id)
        if not target:
            continue
        normalized_target = str(target).replace("\\", "/").lstrip("/")
        if not normalized_target.startswith("xl/"):
            normalized_target = posixpath.normpath(
                posixpath.join("xl", normalized_target)
            )
        if normalized_target in worksheets:
            labels[normalized_target] = str(sheet.attrib.get("name") or Path(normalized_target).stem)
    return labels


def _xlsx_cell_value(cell: Any, shared_strings: list[str]) -> str:
    cell_type = str(cell.attrib.get("t") or "")
    if cell_type == "inlineStr":
        return "".join(
            str(item.text or "")
            for item in cell.iter()
            if _xml_local_name(item.tag) == "t"
        ).strip()
    value = next(
        (
            str(item.text or "")
            for item in cell.iter()
            if _xml_local_name(item.tag) == "v"
        ),
        "",
    ).strip()
    if cell_type == "s":
        try:
            return shared_strings[int(value)].strip()
        except (TypeError, ValueError, IndexError):
            return ""
    return value


def _xlsx_row_number(row: Any, fallback: int) -> int:
    try:
        return max(1, int(row.attrib.get("r") or fallback))
    except (TypeError, ValueError):
        return max(1, fallback)


def _xml_local_name(tag: Any) -> str:
    return str(tag).rsplit("}", 1)[-1]


def _parse_structured(path: Path, kind: str) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8-sig")
    if kind == "json":
        value = json.loads(text)
    else:
        try:
            import yaml

            value = yaml.safe_load(text)
        except ImportError:
            value = _simple_yaml(text)
    if isinstance(value, dict) and isinstance(value.get("metrics"), list):
        value = value["metrics"]
    if isinstance(value, dict):
        value = [value]
    if not isinstance(value, list) or not all(isinstance(item, dict) for item in value):
        raise ValueError(f"{kind} metric source must be an object or list of objects")
    return [{"__source_locator": f"item:{index}", **item} for index, item in enumerate(value, 1)]


def _parse_log(path: Path) -> list[dict[str, Any]]:
    rows = []
    pattern = re.compile(r"(?P<key>[A-Za-z][A-Za-z0-9_.-]*)\s*[:=]\s*(?P<value>-?(?:\d+\.?\d*|\.\d+))")
    for line_number, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
        for match in pattern.finditer(line):
            key = match.group("key")
            lowered = key.casefold()
            if not any(token in lowered for token in ("accuracy", "f1", "precision", "recall", "loss", "latency", "score", "metric")):
                continue
            rows.append({"metric": key, "value": match.group("value"), "__source_locator": f"line:{line_number}"})
    return rows


def _simple_yaml(text: str) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for line in text.splitlines():
        if not line.strip() or line.lstrip().startswith("#") or ":" not in line:
            continue
        key, value = line.split(":", 1)
        value = value.strip().strip("\"'")
        result[key.strip()] = value
    return result


def _normalise_metric_row(raw: dict[str, Any], index: int) -> dict[str, Any]:
    row = {_normalise_key(key): value for key, value in raw.items() if not str(key).startswith("__")}
    return {
        "metric_name": _text(_pick(row, METRIC_ALIASES)),
        "value": _pick(row, VALUE_ALIASES),
        "unit": _text(_pick(row, UNIT_ALIASES)),
        "split": _text(_pick(row, SPLIT_ALIASES)),
        "dataset_version": _text(_pick(row, DATASET_ALIASES)),
        "performed_at": _text(_pick(row, TIME_ALIASES)),
        "direction": _normalise_direction(_pick(row, DIRECTION_ALIASES)),
        "definition_version": _text(_pick(row, DEFINITION_ALIASES)),
        "evaluation_script_ref": _text(_pick(row, SCRIPT_ALIASES)),
        "config_ref": _text(_pick(row, ("config_ref", "configuration"))),
        "environment_ref": _text(_pick(row, ("environment_ref", "environment"))),
        "sample_count": _pick(row, SAMPLE_ALIASES),
        "variance": _pick(row, VARIANCE_ALIASES),
        "source_locator": str(raw.get("__source_locator") or f"row:{index}")[:260],
    }


def _normalise_key(value: Any) -> str:
    return re.sub(r"[\s-]+", "_", str(value or "").strip().casefold())


def _pick(row: dict[str, Any], aliases: tuple[str, ...]) -> Any:
    for alias in aliases:
        key = _normalise_key(alias)
        if key in row and row[key] not in (None, ""):
            return row[key]
    return None


def _text(value: Any) -> str:
    return str(value or "").strip()


def _normalise_direction(value: Any) -> str:
    text = _text(value).lower()
    if text in {"higher", "high", "max", "higher_is_better", "越高越好", "高"}:
        return "higher"
    if text in {"lower", "low", "min", "lower_is_better", "越低越好", "低"}:
        return "lower"
    return text


def _finite_number(value: Any) -> bool:
    try:
        return math.isfinite(float(value))
    except (TypeError, ValueError):
        return False


def _non_negative_int(value: Any, name: str) -> int:
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{name} must be an integer") from exc
    if number < 0:
        raise ValueError(f"{name} must be non-negative")
    return number


def _non_negative_float(value: Any, name: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{name} must be numeric") from exc
    if not math.isfinite(number) or number < 0.0:
        raise ValueError(f"{name} must be finite and non-negative")
    return number


def _parse_timestamp(value: str) -> datetime:
    parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        raise ValueError("timestamp must include a timezone")
    return parsed


def _canonical_timestamp(value: str) -> str:
    return _parse_timestamp(value).astimezone(timezone.utc).isoformat()


def _metric_key(item: dict[str, Any]) -> tuple[str, str, str, str]:
    return (str(item["name"]), str(item["unit"]), str(item["definition_version"]), str(item["split"]))


def _same_non_empty(left: Any, right: Any) -> bool:
    return bool(str(left or "").strip()) and str(left) == str(right)


def _run_ref(run: dict[str, Any]) -> dict[str, Any]:
    return {
        key: run.get(key)
        for key in (
            "run_id", "project_id", "experiment_name", "performed_at", "dataset_version",
            "config_ref", "evaluation_script_ref", "environment_ref", "evidence_version_id",
            "evidence_title", "file_ref",
        )
    }


def _measurement_ref(run: dict[str, Any], measurement: dict[str, Any]) -> dict[str, Any]:
    return {
        "run_id": run["run_id"],
        "evidence_version_id": run.get("evidence_version_id"),
        "file_ref": run.get("file_ref"),
        "source_locator": measurement.get("source_locator") or "未记录",
    }


def _candidate_factors(baseline: dict[str, Any], current: dict[str, Any]) -> list[dict[str, Any]]:
    factors = []
    for field, label in (("config_ref", "配置"), ("environment_ref", "环境")):
        left, right = str(baseline.get(field) or ""), str(current.get(field) or "")
        if left != right:
            factors.append({"factor": field, "label": label, "level": "candidate", "baseline": left, "current": right})
    return factors


def _next_validation(factors: list[dict[str, Any]]) -> str:
    if not factors:
        return "当前账本未记录配置或环境变化；若要解释原因，请补充单因素对照运行。"
    return "固定数据集、评测脚本和其余配置，仅改变候选因素做单因素对照，重跑同一 split，再判断因果。"
